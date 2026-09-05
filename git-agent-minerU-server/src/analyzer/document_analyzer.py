# encoding: utf-8
# @file: document_analyzer.py
# @desc: 文件分析两阶段同步管线。阶段0按扩展名提取内容（优先 MinerU，降级传统解析器），
#        阶段1 LLM 初步分析，阶段2 LLM 结构化字段提取，输出
#        [型号,规格,电压,颜色,标准,数量,单位]。返回 dict 含 code 字段。
import os
from typing import Optional, Dict, Any, List

from analyzer.models import FIELD_MAP, FIELDS_CN

# ---------------- 可选依赖（传统解析器降级用） ----------------
try:
    import pdfplumber
    PDFPLUMBER_AVAILABLE = True
except ImportError:
    PDFPLUMBER_AVAILABLE = False

try:
    import pytesseract
    from PIL import Image
    PYTESSERACT_AVAILABLE = True
except ImportError:
    PYTESSERACT_AVAILABLE = False

try:
    import docx
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False

try:
    import openpyxl
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False


# MinerU 支持解析的扩展名
IMAGE_EXTS = {".jpg", ".jpeg", ".png", ".bmp", ".gif", ".tif", ".tiff", ".webp"}
MINERU_EXTS = {".pdf", ".docx", ".doc", "xls",".xlsx"} | IMAGE_EXTS


def _get_llm(json_mode: bool = False):
    """延迟创建 langchain ChatOpenAI 实例（读取 .env 配置），返回 (llm, model)
    json_mode=True 时启用 response_format=json_object（DeepSeek V4 支持）"""
    from langchain_openai import ChatOpenAI
    from agent.load_env import load_env
    env = load_env()
    kwargs = {
        "model": env["llm_model"],
        "api_key": env["API_KEY"],
        "base_url": env["base_url"],
        "temperature": 0,
    }
    if json_mode:
        kwargs["model_kwargs"] = {"response_format": {"type": "json_object"}}
    return ChatOpenAI(**kwargs), env["llm_model"]


def _parse_json_lenient(text: str):
    """宽松解析 LLM 输出的 JSON：去除代码围栏、提取首个 JSON 对象"""
    import json
    import re
    if not text:
        return None
    text = text.strip()
    if text.startswith("```"):
        text = re.sub(r"^```(?:json)?\s*", "", text)
        text = re.sub(r"\s*```$", "", text).strip()
    try:
        return json.loads(text)
    except Exception:
        m = re.search(r"\{.*\}", text, re.S)
        if m:
            try:
                return json.loads(m.group(0))
            except Exception:
                return None
    return None


class DocumentAnalyzer:
    """
    文件分析器：MinerU 优先 + 传统解析器降级，LLM 两阶段结构化字段提取。
    所有方法返回 JSON 可序列化 dict，成功包含 code=200，失败包含 code 与 error。
    """

    # ================= 传统解析器（降级路径） =================
    @staticmethod
    def extract_text_content(text_path: str) -> Dict[str, Any]:
        if not os.path.exists(text_path):
            return {"error": f"文件 {text_path} 不存在"}
        for encoding in ("utf-8-sig", "gbk"):
            try:
                with open(text_path, "r", encoding=encoding) as f:
                    return {"text": f.read(), "parser": f"plain-text({encoding})"}
            except UnicodeDecodeError:
                continue
            except Exception as e:
                return {"error": f"文本文件读取失败: {e}"}
        return {"error": "无法识别文件编码，请确保为 UTF-8 或 GBK 编码的文本文件"}

    @staticmethod
    def extract_word_content(word_path: str) -> Dict[str, Any]:
        if not os.path.exists(word_path):
            return {"error": f"文件 {word_path} 不存在"}
        if not DOCX_AVAILABLE:
            return {"error": "未安装 Word 解析库，请执行: pip install python-docx"}
        try:
            doc = docx.Document(word_path)
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            tables = []
            for table in doc.tables:
                tables.append([[cell.text for cell in row.cells] for row in table.rows])
            return {"text": "\n".join(paragraphs), "tables": tables, "parser": "python-docx"}
        except Exception as e:
            return {"error": f"Word 解析失败: {e}"}

    @staticmethod
    def extract_excel_content(excel_path: str) -> Dict[str, Any]:
        if not os.path.exists(excel_path):
            return {"error": f"文件 {excel_path} 不存在"}
        if not OPENPYXL_AVAILABLE:
            return {"error": "未安装 Excel 解析库，请执行: pip install openpyxl"}
        try:
            wb = openpyxl.load_workbook(excel_path, data_only=True, read_only=True)
            sheets = {}
            for name in wb.sheetnames:
                rows = []
                for row in wb[name].iter_rows(values_only=True):
                    if any(c is not None and str(c).strip() for c in row):
                        rows.append(list(row))
                sheets[name] = rows
            wb.close()
            return {"sheets": sheets, "parser": "openpyxl"}
        except Exception as e:
            return {"error": f"Excel 解析失败: {e}（旧版 .xls 请先另存为 .xlsx）"}

    @staticmethod
    def extract_image_content(image_path: str) -> Dict[str, Any]:
        if not os.path.exists(image_path):
            return {"error": f"文件 {image_path} 不存在"}
        if not PYTESSERACT_AVAILABLE:
            return {"error": "未安装 OCR 库，请执行: pip install pytesseract pillow 并安装 Tesseract-OCR"}
        try:
            img = Image.open(image_path)
            text = pytesseract.image_to_string(img, lang="chi_sim+eng")
            if text.strip():
                return {"text": text.strip(), "parser": "pytesseract"}
            return {"error": "图片 OCR 未识别到文字内容（图片可能空白或清晰度过低）"}
        except pytesseract.TesseractNotFoundError:
            return {"error": "未安装 Tesseract-OCR 引擎或缺少中文语言包(chi_sim)"}
        except Exception as e:
            return {"error": f"图片 OCR 识别失败: {e}"}

    @staticmethod
    def extract_pdf_content(pdf_path: str) -> Dict[str, Any]:
        if not os.path.exists(pdf_path):
            return {"error": f"文件 {pdf_path} 不存在"}
        if PDFPLUMBER_AVAILABLE:
            text_parts, tables, pages = [], [], 0
            try:
                with pdfplumber.open(pdf_path) as pdf:
                    pages = len(pdf.pages)
                    for page in pdf.pages:
                        t = page.extract_text() or ""
                        if t.strip():
                            text_parts.append(t)
                        for tbl in page.extract_tables():
                            tables.append(tbl)
                return {"text": "\n".join(text_parts), "tables": tables, "pages": pages, "parser": "pdfplumber"}
            except Exception as e:
                return {"error": f"PDF 解析失败: {e}"}
        return {"error": "未安装 PDF 解析库，请执行: pip install pdfplumber"}

    # ================= MinerU 高精度解析 =================
    @staticmethod
    def extract_via_mineru(file_path: str, logger=None) -> Dict[str, Any]:
        """
        调用本地 MinerU 解析文件，返回 {"text": markdown, "parser": "mineru"} 或 {"error": ...}
        仅对 MinerU 支持的类型调用；不支持类型直接返回 error。
        """
        ext = os.path.splitext(file_path)[1].lower()
        if ext not in MINERU_EXTS:
            return {"error": f"MinerU 不支持解析该类型: {ext}"}
        try:
            from mineru_parser.mineru_client import MinerUClient
            client = MinerUClient.get_instance(logger=logger)
            md = client.parse_to_markdown(file_path)
            return {"text": md, "parser": "mineru"}
        except Exception as e:
            # 带上异常类型，避免出现空信息导致无法定位
            detail = str(e).strip() or "(无详情)"
            return {"error": f"MinerU 解析失败[{type(e).__name__}]: {detail}"}

    @staticmethod
    def _render_tables_text(tables: List) -> str:
        if not tables:
            return ""
        lines = []
        for i, table in enumerate(tables, 1):
            lines.append(f"[表格{i}]")
            for row in table:
                cells = [str(c) if c is not None else "" for c in row]
                lines.append(" | ".join(cells))
        return "\n".join(lines)

    # ================= LLM 两阶段分析 =================
    @staticmethod
    def _llm_first_analysis(content: str, question: str) -> str:
        """阶段1: 根据用户分析需求，用 LLM 对文件内容初步分析；无需求返回空串"""
        if not question or not question.strip():
            return ""
        if len(content) > 12000:
            content = content[:12000]
        try:
            from langchain_core.messages import SystemMessage, HumanMessage
            llm, _ = _get_llm()
            resp = llm.invoke([
                SystemMessage(content=(
                    "你是文件内容分析专家。请严格根据用户的分析需求，对给定文件内容进行分析解读，"
                    "输出与需求直接相关的分析结论。只基于文件内容陈述，不编造数据，用中文简洁输出。"
                )),
                HumanMessage(content=f"用户分析需求: {question}\n\n文件内容:\n{content}")
            ])
            return resp.content or ""
        except Exception as e:
            return f"(初步分析失败: {e})"

    @staticmethod
    def _llm_analyze_fields(source: str, content: str, tables_text: str,
                            doc_type: str, parser: str, pages: Optional[int],
                            first_analysis: str = "") -> Dict[str, Any]:
        """阶段2: LLM 将文档内容按字段结构化提取（json_object 模式），返回含 code 的 dict"""
        if len(content) > 30000:
            content = content[:30000]
        try:
            from langchain_core.messages import SystemMessage, HumanMessage
            llm, _ = _get_llm(json_mode=True)
            schema_example = (
                '{"items": [{"型号": "YJV", "规格": "3x2.5mm²", "电压": "0.6/1kV", '
                '"颜色": "红", "标准": "GB/T 12706", "数量": 500, "单位": "米"}]}'
            )
            prompt = (
                "你是电线电缆物料清单解析专家。请从以下文档内容中提取全部物料条目，"
                "每个条目按字段：型号、规格、电压、颜色、标准、数量、单位 提取。\n"
                "要求：\n"
                "1. 按出现顺序输出所有条目。\n"
                "2. 缺失字段用 null，不要编造原文中不存在的数据。\n"
                "3. 数量输出为数字（如原文\"500米\"：数量=500，单位=\"米\"）。\n"
                "4. 电压为额定电压等级（如 0.6/1kV、450/750V）；标准为执行标准编号（如 GB/T 5023）。\n"
                "5. 只输出一个 JSON 对象，格式如下，不要输出任何其它内容或解释：\n"
                f"{schema_example}\n"
            )
            if first_analysis:
                prompt += (
                    "\n请结合下面的初步分析结论进行二次分析校准（与原文冲突时以原文为准）：\n"
                    f"初步分析结论：\n{first_analysis}\n"
                )
            prompt += (
                f"\n文档内容：\n{content}\n\n文档表格：\n{tables_text if tables_text else '（无表格）'}"
            )
            resp = llm.invoke([
                SystemMessage(content="你是电线电缆物料清单解析专家，只输出 JSON 对象。"),
                HumanMessage(content=prompt)
            ])
            raw = resp.content or ""
            data = _parse_json_lenient(raw)
        except Exception as e:
            return {
                "code": 503,
                "error": f"LLM 字段分析失败: {e}",
                "source": source,
                "metadata": {"parser": parser, "pages": pages}
            }

        if not isinstance(data, dict) or "items" not in data:
            return {
                "code": 503,
                "error": "LLM 返回内容无法解析为物料条目",
                "source": source,
                "metadata": {"parser": parser, "pages": pages}
            }
        items = []
        for it in (data.get("items") or []):
            if not isinstance(it, dict):
                continue
            items.append({cn: it.get(cn) for cn in FIELDS_CN})
        return {
            "code": 200,
            "type": doc_type,
            "source": source,
            "content": content,
            "analysis": {
                "fields": FIELDS_CN,
                "items": items,
                "item_count": len(items)
            },
            "metadata": {"parser": parser, "pages": pages}
        }

    # ================= 主管线 =================
    @staticmethod
    def analyze_file_pipeline(file_path: str, question: str = "",
                              logger=None) -> Dict[str, Any]:
        """
        两阶段同步分析管线（按扩展名自动路由）：
          阶段0: 提取内容（优先 MinerU，失败降级传统解析器）
          阶段1: LLM 根据需求初步分析
          阶段2: LLM 结构化字段提取
        返回 JSON 可序列化 dict，含 code（200/415/422/503）、answer、analysis 等。
        """
        ext = os.path.splitext(file_path)[1].lower()
        content = ""
        tables_text = ""
        doc_type = "text"
        parser = "unknown"
        pages = None

        # ----- 阶段0: 提取内容 -----
        # .txt 不支持 MinerU，直接走文本读取
        use_mineru = ext in MINERU_EXTS
        mineru_ok = False
        if use_mineru:
            mu = DocumentAnalyzer.extract_via_mineru(file_path, logger=logger)
            if "error" not in mu:
                content = mu["text"]
                tables_text = ""
                doc_type = ext.lstrip(".") or "file"
                parser = mu["parser"]
                pages = None
                mineru_ok = True
            else:
                if logger:
                    logger(f"MinerU 不可用，降级传统解析: {mu['error']}")

        if not mineru_ok:
            try:
                if ext == ".pdf":
                    ex = DocumentAnalyzer.extract_pdf_content(file_path)
                    if "error" in ex:
                        return {"code": 422, "error": ex["error"], "source": file_path}
                    content = ex["text"]
                    tables_text = DocumentAnalyzer._render_tables_text(ex.get("tables") or [])
                    doc_type, parser, pages = "pdf", ex["parser"], ex["pages"]
                elif ext == ".docx" or ext == ".doc":
                    ex = DocumentAnalyzer.extract_word_content(file_path)
                    if "error" in ex:
                        return {"code": 422, "error": ex["error"], "source": file_path}
                    content = ex["text"]
                    tables_text = DocumentAnalyzer._render_tables_text(ex.get("tables") or [])
                    doc_type, parser, pages = "word", ex["parser"], None
                elif ext in (".xlsx", ".xlsm"):
                    ex = DocumentAnalyzer.extract_excel_content(file_path)
                    if "error" in ex:
                        return {"code": 422, "error": ex["error"], "source": file_path}
                    lines = []
                    for name, rows in ex["sheets"].items():
                        lines.append(f"[工作表: {name}]")
                        for row in rows:
                            lines.append(" | ".join("" if c is None else str(c) for c in row))
                    content = "\n".join(lines)
                    tables_text = ""
                    doc_type, parser, pages = "excel", ex["parser"], None
                elif ext in IMAGE_EXTS:
                    ex = DocumentAnalyzer.extract_image_content(file_path)
                    if "error" in ex:
                        return {"code": 422, "error": ex["error"], "source": file_path}
                    content = ex["text"]
                    tables_text = ""
                    doc_type, parser, pages = "image", ex["parser"], None
                elif ext in (".txt", ".csv", ".md", ".log"):
                    ex = DocumentAnalyzer.extract_text_content(file_path)
                    if "error" in ex:
                        return {"code": 422, "error": ex["error"], "source": file_path}
                    content = ex["text"]
                    tables_text = ""
                    doc_type, parser, pages = "text", ex["parser"], None
                else:
                    return {"code": 415, "error": f"不支持的文件类型: {ext}", "source": file_path}
            except Exception as e:
                return {"code": 422, "error": f"文件提取失败: {e}", "source": file_path}

        if not content.strip() and not tables_text.strip():
            return {
                "code": 422,
                "error": "未提取到文本内容（可能为扫描件/图片不清晰，建议启用 MinerU 高精度解析）",
                "source": file_path
            }

        # ----- 阶段1: LLM 初步分析 -----
        first_analysis = DocumentAnalyzer._llm_first_analysis(content, question)

        # ----- 阶段2: LLM 结构化字段提取 -----
        result = DocumentAnalyzer._llm_analyze_fields(
            source=file_path, content=content, tables_text=tables_text,
            doc_type=doc_type, parser=parser, pages=pages,
            first_analysis=first_analysis
        )
        if result.get("code") != 200:
            return result

        items = result["analysis"]["items"]
        if first_analysis and not first_analysis.startswith("(初步分析失败"):
            answer = f"{first_analysis}\n\n【字段提取结果】共识别 {len(items)} 条物料条目"
        else:
            answer = f"已完成字段分析，共识别 {len(items)} 条物料条目"
            if first_analysis:
                answer += f"（需求初步分析: {first_analysis}）"

        result["first_analysis"] = first_analysis
        result["answer"] = answer
        return result
